#!/usr/bin/env python3
"""Build a Sankey budget diagram PNG from a Nextcloud budget workbook.

Runs as a systemd oneshot triggered by a timer (see
sankey-export.service/.timer). Reads the workbook and writes the PNG over
WebDAV only -- it never touches Nextcloud's data directory on disk. A
cheap WebDAV ETag check skips the expensive render entirely when the
workbook hasn't changed since the last run, which is what makes running
this every minute cheap rather than wasteful.

The workbook itself is still Finanzfluss' own free budget template
(same "Einnahmen"/"Ausgaben" sheet layout as ever -- that's just a
downloadable spreadsheet, not a live dependency), but the diagram is
rendered entirely locally with Plotly/Kaleido now, not by driving a
browser against finanzfluss.de and scraping their own Highcharts export.
That scrape was the actual fragility this replaced: it broke the moment
their site changed its export flow, rate-limited harder, or added bot
detection, with zero warning and nothing this repo could do about it.
Local rendering has no dependency on any third party being up, unchanged,
or in a good mood.
"""

from __future__ import annotations

import math
import os
import sys
import xml.etree.ElementTree as ET
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
from openpyxl import load_workbook

DAV = "DAV:"

ENDPOINT_HOST = os.environ.get("SANKEY_EXPORT_ENDPOINT_HOST", "127.0.0.1")
ENDPOINT_PORT = int(os.environ.get("SANKEY_EXPORT_ENDPOINT_PORT", "11000"))
HTTP_HOST = os.environ.get("SANKEY_EXPORT_HTTP_HOST", "nextcloud.jkandler.de")
USERNAME = os.environ.get("SANKEY_EXPORT_USERNAME", "")
APP_PASSWORD_FILE = os.environ.get(
    "SANKEY_EXPORT_APP_PASSWORD_FILE", "/etc/sankey-export/app-password"
)
REMOTE_DIR = os.environ.get("SANKEY_EXPORT_REMOTE_DIR", "Documents/Finanzen").strip("/")
WORKBOOK_NAME = os.environ.get("SANKEY_EXPORT_WORKBOOK_NAME", "Finanzfluss_Nextcloud.xlsx")
OUTPUT_DIR = Path(os.environ.get("SANKEY_EXPORT_OUTPUT_DIR", "/var/lib/sankey-export/out"))
STATE_FILE = Path(
    os.environ.get("SANKEY_EXPORT_STATE_FILE", "/var/lib/sankey-export/last-etag")
)

# Node colors -- income green, the hub a neutral blue-grey, "Budget"
# (money left over) gold, expense categories cycling through a
# categorical palette. Subcategories inherit their own category's color.
INCOME_COLOR = "#2e7d32"
HUB_COLOR = "#455a64"
BUDGET_COLOR = "#f9a825"
CATEGORY_PALETTE = [
    "#c62828", "#6a1b9a", "#1565c0", "#00838f", "#ad1457",
    "#4527a0", "#00695c", "#e65100", "#283593", "#558b2f",
]


@dataclass(frozen=True)
class IncomeRow:
    name: str
    amount: float


@dataclass(frozen=True)
class CostRow:
    category: str
    subcategory: str
    amount: float


@dataclass(frozen=True)
class SankeySubcategory:
    name: str
    amount: float


@dataclass(frozen=True)
class SankeyCategory:
    name: str
    amount: float
    percentage: float
    color: str
    subcategories: list[SankeySubcategory] = field(default_factory=list)


@dataclass(frozen=True)
class SankeyIncome:
    name: str
    amount: float
    percentage: float


@dataclass(frozen=True)
class BudgetSummary:
    """Pure, renderer-independent summary of one workbook's budget.

    Deliberately has no Plotly/Kaleido dependency at all -- this is the
    part worth unit-testing directly (category grouping, subcategory
    dedup, the negative-budget warning), the same way this repo's own
    Terraform/Ansible split keeps "compute the right thing" separate
    from "actually apply it".
    """

    incomes: list[SankeyIncome]
    categories: list[SankeyCategory]
    income_total: float
    expense_total: float
    budget: float
    warnings: list[str]


def read_app_password(path: str) -> str:
    value = Path(path).read_text(encoding="utf-8").strip()
    if len(value) < 16 or value == "CHANGE_ME":
        raise RuntimeError("Nextcloud app password is not configured")
    return value


def active(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().lower() in {"1", "true", "wahr", "yes", "ja", "x", "✓", "✅"}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def clean_amount(value: Any, *, row_description: str) -> float:
    if value is None or value == "":
        raise ValueError(f"Missing amount for {row_description}")
    if isinstance(value, str):
        normalized = value.strip().replace("€", "").replace(" ", "")
        if "," in normalized and "." in normalized:
            normalized = normalized.replace(".", "").replace(",", ".")
        else:
            normalized = normalized.replace(",", ".")
        value = normalized
    try:
        amount = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid amount for {row_description}: {value!r}") from exc
    if not math.isfinite(amount) or amount <= 0:
        raise ValueError(f"Amount must be greater than 0 for {row_description}: {amount}")
    return round(amount, 2)


def format_euro(amount: float) -> str:
    """German-style formatting (1.234,56 €) -- matches the workbook's own locale."""
    return f"{amount:,.2f} €".replace(",", "¬").replace(".", ",").replace("¬", ".")


def percentage(amount: float, total: float) -> float:
    return amount / total * 100 if total else 0.0


# The workbook's sheet/tab names ("Einnahmen"/"Ausgaben") and its budget
# row name ("Budget") are the real template's literal structure, not our
# own UI text -- they must match the actual spreadsheet, not be translated.
def read_workbook(workbook_path: Path) -> tuple[list[IncomeRow], list[CostRow]]:
    wb = load_workbook(workbook_path, data_only=False)
    missing = {"Einnahmen", "Ausgaben"} - set(wb.sheetnames)
    if missing:
        raise ValueError(f"Missing worksheet(s): {', '.join(sorted(missing))}")

    income_sheet = wb["Einnahmen"]
    cost_sheet = wb["Ausgaben"]

    incomes: list[IncomeRow] = []
    for row in range(4, income_sheet.max_row + 1):
        if not active(income_sheet.cell(row, 1).value):
            continue
        name = clean_text(income_sheet.cell(row, 2).value)
        amount_value = income_sheet.cell(row, 3).value
        if not name and (amount_value is None or amount_value == ""):
            continue
        if not name:
            raise ValueError(f"'Einnahmen' sheet, row {row}: name is missing")
        amount = clean_amount(amount_value, row_description=f"'Einnahmen' row {row} ({name})")
        incomes.append(IncomeRow(name=name, amount=amount))

    costs: list[CostRow] = []
    for row in range(4, cost_sheet.max_row + 1):
        if not active(cost_sheet.cell(row, 1).value):
            continue
        category = clean_text(cost_sheet.cell(row, 2).value)
        subcategory = clean_text(cost_sheet.cell(row, 3).value)
        amount_value = cost_sheet.cell(row, 4).value
        if not category and not subcategory and (amount_value is None or amount_value == ""):
            continue
        if not category:
            raise ValueError(f"'Ausgaben' sheet, row {row}: category is missing")

        # "Budget" is calculated automatically from income minus real expenses.
        # A legacy Budget row in the workbook is therefore ignored completely.
        if category.casefold() == "budget":
            continue

        amount = clean_amount(amount_value, row_description=f"'Ausgaben' row {row} ({category})")
        costs.append(CostRow(category=category, subcategory=subcategory, amount=amount))

    if not incomes:
        raise ValueError("No active income rows found")
    if not costs:
        raise ValueError("No active expense rows found")
    return incomes, costs


def summarize_budget(incomes: list[IncomeRow], costs: list[CostRow]) -> BudgetSummary:
    income_total = round(sum(row.amount for row in incomes), 2)
    expense_total = round(sum(row.amount for row in costs), 2)
    budget = round(income_total - expense_total, 2)
    warnings: list[str] = []

    income_summaries = [
        SankeyIncome(name=row.name, amount=row.amount, percentage=percentage(row.amount, income_total))
        for row in incomes
    ]

    grouped: "OrderedDict[str, list[CostRow]]" = OrderedDict()
    for row in costs:
        grouped.setdefault(row.category, []).append(row)

    categories: list[SankeyCategory] = []
    for index, (category, rows) in enumerate(grouped.items()):
        color = CATEGORY_PALETTE[index % len(CATEGORY_PALETTE)]
        category_total = round(sum(row.amount for row in rows), 2)
        subcategories: list[SankeySubcategory] = []
        has_named_subcategory = any(row.subcategory for row in rows)

        if has_named_subcategory:
            used_names: set[str] = set()
            for row in rows:
                # "Sonstiges" ("Miscellaneous") stays German: it becomes a
                # real label on the generated Sankey diagram, alongside the
                # user's own German category names -- translating just this
                # one fallback would look inconsistent on the chart itself.
                sub_name = row.subcategory or "Sonstiges"
                original_name = sub_name
                suffix = 2
                while sub_name in used_names:
                    sub_name = f"{original_name} {suffix}"
                    suffix += 1
                used_names.add(sub_name)

                if not row.subcategory:
                    warnings.append(
                        f"Category '{category}': blank subcategory was "
                        "exported as 'Sonstiges'."
                    )
                subcategories.append(SankeySubcategory(name=sub_name, amount=row.amount))

        categories.append(
            SankeyCategory(
                name=category,
                amount=category_total,
                percentage=percentage(category_total, income_total),
                color=color,
                subcategories=subcategories,
            )
        )

    if budget < 0:
        warnings.append(
            f"Expenses exceed income by {abs(budget):.2f} €. A negative "
            "budget can't be shown on the Sankey diagram as a normal "
            "expense category."
        )

    return BudgetSummary(
        incomes=income_summaries,
        categories=categories,
        income_total=income_total,
        expense_total=expense_total,
        budget=budget,
        warnings=warnings,
    )


def build_sankey_figure(summary: BudgetSummary):
    """Build the Plotly Figure from a BudgetSummary. Pure/no I/O -- the
    actual PNG write (and the only part that needs Kaleido/a real
    browser) lives in render_sankey_png below.
    """
    import plotly.graph_objects as go

    labels: list[str] = []
    colors: list[str] = []
    sources: list[int] = []
    targets: list[int] = []
    values: list[float] = []

    def add_node(label: str, color: str) -> int:
        labels.append(label)
        colors.append(color)
        return len(labels) - 1

    hub = add_node(f"Einkommen<br>{format_euro(summary.income_total)}", HUB_COLOR)

    for income in summary.incomes:
        node = add_node(
            f"{income.name}<br>{format_euro(income.amount)} ({income.percentage:.0f}%)",
            INCOME_COLOR,
        )
        sources.append(node)
        targets.append(hub)
        values.append(income.amount)

    for category in summary.categories:
        category_node = add_node(
            f"{category.name}<br>{format_euro(category.amount)} ({category.percentage:.0f}%)",
            category.color,
        )
        sources.append(hub)
        targets.append(category_node)
        values.append(category.amount)

        for subcategory in category.subcategories:
            sub_node = add_node(
                f"{subcategory.name}<br>{format_euro(subcategory.amount)}",
                category.color,
            )
            sources.append(category_node)
            targets.append(sub_node)
            values.append(subcategory.amount)

    if summary.budget > 0:
        budget_node = add_node(
            f"Budget<br>{format_euro(summary.budget)} "
            f"({percentage(summary.budget, summary.income_total):.0f}%)",
            BUDGET_COLOR,
        )
        sources.append(hub)
        targets.append(budget_node)
        values.append(summary.budget)

    fig = go.Figure(
        go.Sankey(
            arrangement="snap",
            node=dict(label=labels, color=colors, pad=18, thickness=16, line=dict(width=0)),
            link=dict(source=sources, target=targets, value=values, color="rgba(160,160,160,0.35)"),
        )
    )
    fig.update_layout(
        title=(
            f"Budget {format_euro(summary.income_total)} Einkommen / "
            f"{format_euro(summary.expense_total)} Ausgaben"
        ),
        font_size=13,
        width=1600,
        height=1000,
        margin=dict(l=10, r=10, t=60, b=10),
    )
    return fig


def render_sankey_png(summary: BudgetSummary, output_path: Path) -> None:
    # scale=2 for a sharp, roughly-HiDPI PNG -- matches the resolution
    # bump the old Highcharts export applied via its own "scale" param.
    build_sankey_figure(summary).write_image(str(output_path), scale=2)


class NextcloudWebDAV:
    """WebDAV client fixed to the loopback AIO endpoint and one account."""

    def __init__(self, host: str, port: int, http_host: str, username: str, app_password: str) -> None:
        self.base_url = f"http://{host}:{port}"
        self.http_host = http_host
        self.username = username
        self.root = f"{self.base_url}/remote.php/dav/files/{quote(username, safe='')}"
        self.session = requests.Session()
        self.session.auth = (username, app_password)
        self.session.headers.update(
            {"User-Agent": "sankey-export/1", "Host": http_host}
        )

    def _url(self, remote_path: str) -> str:
        parts = [quote(part, safe="") for part in remote_path.strip("/").split("/") if part]
        return self.root + ("/" + "/".join(parts) if parts else "")

    def etag(self, remote_path: str) -> str | None:
        """Cheap PROPFIND for just the ETag -- no content transferred."""
        body = b'<?xml version="1.0"?><d:propfind xmlns:d="DAV:"><d:prop><d:getetag/></d:prop></d:propfind>'
        response = self.session.request(
            "PROPFIND",
            self._url(remote_path),
            data=body,
            headers={"Content-Type": "application/xml", "Depth": "0"},
            timeout=30,
        )
        if response.status_code == 404:
            return None
        if response.status_code != 207:
            raise RuntimeError(f"ETag lookup failed ({response.status_code}): {remote_path}")
        root = ET.fromstring(response.content)
        etag_element = root.find(f".//{{{DAV}}}getetag")
        return etag_element.text if etag_element is not None else None

    def ensure_dir(self, remote_dir: str) -> None:
        current = ""
        for part in [p for p in remote_dir.strip("/").split("/") if p]:
            current = f"{current}/{part}"
            response = self.session.request("MKCOL", self._url(current), timeout=30)
            if response.status_code not in {201, 405}:
                raise RuntimeError(f"Could not create Nextcloud folder ({response.status_code}): {current}")

    def download(self, remote_path: str, local_path: Path) -> None:
        response = self.session.get(self._url(remote_path), timeout=120)
        if response.status_code != 200:
            raise RuntimeError(f"Download failed ({response.status_code}): {remote_path}\n{response.text[:300]}")
        local_path.write_bytes(response.content)

    def upload(self, local_path: Path, remote_path: str) -> None:
        remote_dir = remote_path.rsplit("/", 1)[0] if "/" in remote_path.strip("/") else ""
        if remote_dir:
            self.ensure_dir(remote_dir)
        with local_path.open("rb") as handle:
            response = self.session.put(self._url(remote_path), data=handle, timeout=180)
        if response.status_code not in {200, 201, 204}:
            raise RuntimeError(f"Upload failed ({response.status_code}): {remote_path}\n{response.text[:300]}")


def load_cached_etag(state_path: Path) -> str | None:
    try:
        return state_path.read_text(encoding="utf-8").strip() or None
    except OSError:
        return None


def save_cached_etag(state_path: Path, etag: str) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(etag + "\n", encoding="utf-8")


def main() -> int:
    remote_workbook = f"{REMOTE_DIR}/{WORKBOOK_NAME}"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    webdav = NextcloudWebDAV(
        ENDPOINT_HOST, ENDPOINT_PORT, HTTP_HOST, USERNAME, read_app_password(APP_PASSWORD_FILE)
    )

    etag = webdav.etag(remote_workbook)
    if etag is None:
        print(f"Workbook not found: {remote_workbook}", file=sys.stderr)
        return 1

    if etag == load_cached_etag(STATE_FILE):
        print("Workbook unchanged, nothing to do.")
        return 0

    print(f"Workbook changed (ETag {etag}), downloading {remote_workbook} …")
    temporary_workbook = OUTPUT_DIR / f".{Path(WORKBOOK_NAME).stem}.download{Path(WORKBOOK_NAME).suffix}"
    try:
        webdav.download(remote_workbook, temporary_workbook)
        incomes, costs = read_workbook(temporary_workbook)
    finally:
        temporary_workbook.unlink(missing_ok=True)

    summary = summarize_budget(incomes, costs)

    # Filename deliberately unchanged from the old Finanzfluss-scraping
    # version -- an existing Nextcloud share link/embed pointed at this
    # exact path, and renaming it would silently break that for a
    # rendering-internals change that has nothing to do with the
    # filename.
    png_path = OUTPUT_DIR / "finanzfluss-diagramm.png"
    png_path.unlink(missing_ok=True)
    render_sankey_png(summary, png_path)

    if png_path.exists():
        remote_png = f"{REMOTE_DIR}/{png_path.name}"
        print(f"Uploading {png_path.name} to Nextcloud …")
        webdav.upload(png_path, remote_png)
        # Only remember this ETag once a fresh PNG has actually been
        # uploaded -- a failed export leaves the cache stale so the next
        # (1-minute-later) tick retries automatically.
        save_cached_etag(STATE_FILE, etag)

    print(f"Income:   {summary.income_total:.2f} €")
    print(f"Expenses: {summary.expense_total:.2f} €")
    print(f"Budget:   {summary.budget:.2f} €")

    if summary.warnings:
        print("Warnings:")
        for warning in summary.warnings:
            print(f"- {warning}")

    return 0 if png_path.exists() else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        raise SystemExit(130)
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1)
